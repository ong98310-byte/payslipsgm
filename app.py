import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import os
import zipfile
import tempfile

# 网页的标题和说明
st.set_page_config(page_title="J&T 批量工资单生成器", page_icon="📦")
st.title("📦 J&T 批量 Payslip 一键生成系统")
st.markdown("**专属内部工具**：只要把算好工资的 Excel 丢进来，自动排版并按名字打包下载！")

# 创建上传框
uploaded_file = st.file_uploader("📂 拖拽或点击上传你的工资总表 (Excel格式)", type=["xlsx"])

if uploaded_file is not None:
    if st.button("🚀 一键批量生成"):
        with st.spinner('正在拼命排版生成中，请稍候...'):
            try:
                # 读取数据
                df = pd.read_excel(uploaded_file)
                
                # 创建临时文件夹来存放生成的文件
                temp_dir = tempfile.mkdtemp()
                output_folder = os.path.join(temp_dir, "Payslips")
                os.makedirs(output_folder)
                
                # 样式设定
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="2B2B2B", end_color="2B2B2B", fill_type="solid")
                thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                
                # 开始循环生成每个人薪水单
                for index, row in df.iterrows():
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Payslip"
                    
                    # 按照你之前的格式构建排版 (此处套用 G.M 内部结算排版)
                    payslip_data = [
                        ["PRIVATE & CONFIDENTIAL", "", "", ""],
                        ["COMPANY NAME: GOLD MOUNTAIN EXPRESS SDN. BHD.", "", "", ""],
                        ["ROC NO: 1526589-H", "", "", ""],
                        ["", "", "", ""],
                        ["Employee name", f": {row['Name']}", "Payment Month:", f"{row['Month']}"],
                        ["IC number", f": {row['IC']}", "Employee ID:", f"{row['ID']}"],
                        ["Bank Name", f": {row['Bank']}", "Designation:", f"{row['Designation']}"],
                        ["Bank Account Number", f": {row['Account']}", "Department:", f"{row['Dept']}"],
                        ["", "", "", ""],
                        ["Earnings", "Amount (MYR)", "Deductions", "Amount (MYR)"],
                        ["(+) Commision", row['Commision'], "(-) Inactive", row['Inactive']],
                        ["(+) Transit", row['Transit'], "(-) Fake Sign Penalty", row['Fake_Sign_Penalty']],
                        ["(+) Yoyi", row['Yoyi'], "(-) COD Late Penalty", row['COD_Late_Penalty']],
                        ["(+) TTR Pickup", row['TTR_Pickup'], "(-) COD Issue", row['COD_Issue']],
                        ["(+) Duitnow Refund", row['Duitnow_Refund'], "(-) Lost Parcel", row['Lost_Parcel']],
                        ["(+) Attendance", row['Attendance'], "(-) Lateness", row['Lateness']],
                        ["(+) Route Leader", row['Route_Leader'], "(-) PDNC", row['PDNC']],
                        ["(+) Reward", row['Reward'], "(-) Others", row['Others']],
                        ["", "", "", ""],
                        ["Total Earning (MYR)", row['Total_Earning'], "Total Deduction (MYR)", row['Total_Deduction']],
                        ["", "", "Net Pay (MYR)", row['Net_Pay']],
                        ["", "", "TOTAL PARCEL", row['Total_Parcel']],
                        ["", "", "", ""],
                        ["Remark: If you have any questions regarding your salary, please contact the head of department or team leader directly.", "", "", ""]
                    ]
                    
                    for r_idx, r_data in enumerate(payslip_data, 1):
                        ws.append(r_data)
                        for c_idx, cell_value in enumerate(r_data, 1):
                            cell = ws.cell(row=r_idx, column=c_idx)
                            if r_idx == 10: 
                                cell.fill = header_fill
                                cell.font = header_font
                            elif str(cell_value) != "" and r_idx > 10 and r_idx <= 22:
                                cell.border = thin_border
                            if r_idx == 21 and c_idx >= 3:
                                cell.font = Font(bold=True, color="E3000F")

                    ws.column_dimensions['A'].width = 25
                    ws.column_dimensions['B'].width = 15
                    ws.column_dimensions['C'].width = 25
                    ws.column_dimensions['D'].width = 15
                    
                    # 命名文件并保存
                    file_name = f"{row['ID']} {row['Name']} {row['Month']} Payslip.xlsx".replace("/", "_")
                    wb.save(os.path.join(output_folder, file_name))

                # 将所有生成的 Excel 打包成 ZIP
                zip_path = os.path.join(temp_dir, "JnT_Payslips_Batch.zip")
                with zipfile.ZipFile(zip_path, 'w') as zipf:
                    for root, dirs, files in os.walk(output_folder):
                        for file in files:
                            zipf.write(os.path.join(root, file), file)
                
                # 提供下载按钮
                with open(zip_path, "rb") as fp:
                    st.success(f"✅ 成功生成 {len(df)} 份 Payslip！")
                    st.download_button(
                        label="📥 点击下载打包文件 (ZIP)",
                        data=fp,
                        file_name="JnT_Payslips_Batch.zip",
                        mime="application/zip"
                    )
            except Exception as e:
                st.error(f"生成时遇到错误，请检查你的 Excel 表头名称是否正确。详细错误: {e}")